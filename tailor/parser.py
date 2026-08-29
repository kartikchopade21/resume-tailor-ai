"""Resume ingestion: raw file -> text -> structured ResumeDoc + format audit."""
from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field

from .llm import complete_json
from .schema import ResumeDoc


@dataclass
class FormatAudit:
    """Things that break real ATS parsers, detected from the source file."""

    has_tables: bool = False
    has_images: bool = False
    has_text_boxes: bool = False
    multi_column: bool = False
    in_header_footer: bool = False
    non_standard_bullets: bool = False
    unusual_fonts: list[str] = field(default_factory=list)
    page_count: int = 1
    word_count: int = 0
    issues: list[str] = field(default_factory=list)

    def add(self, msg: str) -> None:
        if msg not in self.issues:
            self.issues.append(msg)


STANDARD_FONTS = {
    "arial", "calibri", "helvetica", "times new roman", "times", "georgia",
    "garamond", "cambria", "verdana", "tahoma", "book antiqua", "liberation serif",
    "liberation sans", "carlito", "dejavu sans", "roboto", "lato", "open sans",
    "nimbusroman", "nimbussans", "abcdee+calibri",
}

WEIRD_BULLETS = "➢➤◆★✦✔✓»❖▪▶"


# --------------------------------------------------------------------- text
def extract_text(path: str) -> tuple[str, str, FormatAudit]:
    """Return (text, source_format, audit)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return (*_extract_pdf(path), )
    if ext in (".docx", ".doc"):
        return (*_extract_docx(path), )
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        text = fh.read()
    audit = FormatAudit(word_count=len(text.split()))
    _audit_text(text, audit)
    return text, "txt", audit


def _extract_pdf(path: str) -> tuple[str, str, FormatAudit]:
    import pdfplumber

    audit = FormatAudit()
    chunks: list[str] = []
    fonts: set[str] = set()
    with pdfplumber.open(path) as pdf:
        audit.page_count = len(pdf.pages)
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
            if page.find_tables():
                audit.has_tables = True
            if page.images:
                audit.has_images = True
            for ch in page.chars[:4000]:
                fonts.add(str(ch.get("fontname", "")).lower())
            if _looks_multi_column(page):
                audit.multi_column = True

    text = "\n".join(chunks)
    audit.word_count = len(text.split())
    unusual = [
        f for f in fonts
        if f and not any(std in f.replace("-", "").replace(",", " ").lower()
                         for std in STANDARD_FONTS)
    ]
    audit.unusual_fonts = sorted(unusual)[:6]

    if audit.has_tables:
        audit.add("Table layout detected — many ATS parsers read tables out of order.")
    if audit.has_images:
        audit.add("Images/graphics detected — any text inside them is invisible to an ATS.")
    if audit.multi_column:
        audit.add("Multi-column layout detected — column text is often interleaved on parse.")
    if audit.unusual_fonts:
        audit.add(f"Non-standard fonts: {', '.join(audit.unusual_fonts[:3])}.")
    if audit.page_count > 2:
        audit.add(f"{audit.page_count} pages — most recruiters expect 1–2.")
    _audit_text(text, audit)
    return text, "pdf", audit


def _looks_multi_column(page) -> bool:
    """Heuristic: a wide vertical gutter with text on both sides."""
    words = page.extract_words() or []
    if len(words) < 40:
        return False
    width = page.width or 612
    left = sum(1 for w in words if w["x1"] < width * 0.46)
    right = sum(1 for w in words if w["x0"] > width * 0.54)
    middle = sum(1 for w in words if w["x0"] <= width * 0.54 <= w["x1"])
    return left > 20 and right > 20 and middle < max(3, len(words) * 0.03)


def _extract_docx(path: str) -> tuple[str, str, FormatAudit]:
    import docx

    audit = FormatAudit()
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]

    if document.tables:
        audit.has_tables = True
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

    xml = document.element.xml
    if "<w:drawing" in xml or "<w:pict" in xml:
        audit.has_images = True
    if "txbxContent" in xml:
        audit.has_text_boxes = True

    for section in document.sections:
        for para in list(section.header.paragraphs) + list(section.footer.paragraphs):
            if para.text.strip():
                audit.in_header_footer = True
                parts.append(para.text)
        if getattr(section, "_sectPr", None) is not None:
            cols = section._sectPr.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols"
            )
            if cols is not None and int(cols.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num", "1"
            ) or 1) > 1:
                audit.multi_column = True

    fonts = {
        run.font.name.lower()
        for p in document.paragraphs for run in p.runs
        if run.font is not None and run.font.name
    }
    audit.unusual_fonts = sorted(f for f in fonts if f not in STANDARD_FONTS)[:6]

    text = "\n".join(t for t in parts if t and t.strip())
    audit.word_count = len(text.split())

    if audit.has_tables:
        audit.add("Table layout detected — many ATS parsers read tables out of order.")
    if audit.has_images:
        audit.add("Images/graphics detected — text inside them is invisible to an ATS.")
    if audit.has_text_boxes:
        audit.add("Text boxes detected — their content is frequently dropped entirely.")
    if audit.in_header_footer:
        audit.add("Content in the header/footer — commonly ignored by ATS parsers.")
    if audit.multi_column:
        audit.add("Multi-column section layout detected.")
    if audit.unusual_fonts:
        audit.add(f"Non-standard fonts: {', '.join(audit.unusual_fonts[:3])}.")
    _audit_text(text, audit)
    return text, "docx", audit


def _audit_text(text: str, audit: FormatAudit) -> None:
    if any(ch in text for ch in WEIRD_BULLETS):
        audit.non_standard_bullets = True
        audit.add("Decorative bullet glyphs — use a plain bullet or hyphen.")
    if audit.word_count and audit.word_count < 200:
        audit.add(
            f"Only {audit.word_count} words of extractable text — the file may be "
            "image-based or largely unparseable."
        )


# ---------------------------------------------------------------- structure
_PARSE_SYSTEM = """You convert resume text into structured JSON. You are a parser, \
not an editor: never invent, embellish, summarise away, or correct anything. \
Every string you emit must be traceable to the input text.

Return ONLY a JSON object with exactly this shape:
{
  "contact": {"name","email","phone","location","linkedin","github","portfolio","headline"},
  "summary": "the professional summary/objective verbatim, or \\"\\"",
  "skills": {"<Category>": ["skill", ...]},
  "experience": [{"id","company","title","location","dates","bullets":["..."]}],
  "projects": [{"id","name","subtitle","tech":["..."],"bullets":["..."],"link"}],
  "education": [{"institution","degree","dates","details"}],
  "certifications": ["..."],
  "extras": {"<Section name>": ["line", ...]}
}

Rules:
- ids: "exp1","exp2",... and "proj1","proj2",... in document order.
- If the resume has no explicit skill categories, use one category named "Technical Skills".
- "tech" for a project = the technologies named in its title line, tech-stack line, or bullets.
- "subtitle" = the project's role/context/date line if present, else "".
- Put any section you cannot map (awards, publications, volunteering, languages) into "extras".
- Missing fields become "" or []. Never use null.

Named work inside a job is a project. Many resumes list a role and then group its \
bullets under named sub-headings ("Delta Comparison Tool (Python + Jenkins)", \
"AIKEY (Voice Assistant)"). Each such group is a separate entry in "projects":
- "name" = the sub-heading, minus any parenthesised tech list.
- "subtitle" = "<job title> @ <company>" of the role it sits under.
- "tech" = technologies from the parenthesised list and the group's bullets.
- "bullets" = exactly that group's bullets.
Those bullets belong to the project only — do not repeat them on the role. The role \
keeps just its own ungrouped bullets, and an empty "bullets" list is fine.

Completeness is the whole job: every bullet in the input must appear exactly once \
in your output, verbatim. Never merge two bullets, drop one, or shorten a skills \
list. If you are unsure where something belongs, keep it rather than lose it."""


def structure_resume(text: str, source_format: str = "", audit: FormatAudit | None = None) -> ResumeDoc:
    data = complete_json(
        _PARSE_SYSTEM,
        f"Resume text:\n<resume>\n{text.strip()}\n</resume>",
        max_tokens=8000,
        temperature=0.0,
    )
    doc = ResumeDoc.from_dict(data)
    doc.raw_text = text
    doc.source_format = source_format
    if not doc.projects and not doc.experience:
        raise ValueError(
            "No projects or work experience were found in the resume. "
            "Check that the file contains selectable text (not a scanned image)."
        )
    return doc


def load_resume(path: str) -> tuple[ResumeDoc, FormatAudit]:
    text, fmt, audit = extract_text(path)
    if len(text.split()) < 40:
        raise ValueError(
            "Almost no text could be extracted — the file is likely a scanned image. "
            "An ATS would see the same nothing. Re-export it as a text-based PDF or DOCX."
        )
    return structure_resume(text, fmt, audit), audit


def load_resume_bytes(data: bytes, filename: str) -> tuple[ResumeDoc, FormatAudit]:
    """Streamlit path: write the upload to a temp file, then parse."""
    import tempfile

    suffix = os.path.splitext(filename)[1].lower() or ".txt"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        return load_resume(tmp_path)
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
